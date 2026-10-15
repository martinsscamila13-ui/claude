"""Simula a execução do prompt mestre para duas alunas fictícias e gera o .xlsx
exatamente no formato pedido na Parte 5, validando as regras da Parte 3."""
import datetime as dt, sys, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

D = lambda d, m: dt.date(2026, m, d)
HOJE, FERIADO = D(16, 10), D(2, 11)
MARCO = {2: D(25,10), 3: D(25,10), 4: D(25,10), 5: D(1,11), 6: D(1,11), 7: D(1,11),
         8: D(8,11), 9: D(8,11), 10: D(8,11), 11: D(11,11), 1: D(11,11)}
DIAS = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]

TAREFAS = {
 2: [("Redigir missão, visão e valores", 1.5), ("Descrever o problema que o negócio resolve", 1), ("Redigir o diferencial competitivo", 1)],
 3: [("Definir público-alvo e criar a persona", 1.5), ("Listar 3 concorrentes e comparar", 1.5), ("Visitar 1 farmácia como cliente oculto", 2), ("Levantar 3 tendências do mercado", 1), ("Redigir a seção de mercado", 2)],
 4: [("Levantar a legislação sanitária aplicável", 1.5), ("Listar licenças, AFE e exigências do RT", 1.5), ("Redigir a seção legal", 1.5)],
 7: [("Definir o posicionamento", 1), ("Escolher canais e ações de divulgação", 1.5), ("Desenhar a estratégia de fidelização", 1), ("Redigir o plano de marketing", 1.5)],
 8: [("Estimar o investimento inicial", 2), ("Levantar custos fixos mensais", 1.5), ("Levantar custos variáveis", 1), ("Estimar faturamento mensal", 1.5), ("Calcular ponto de equilíbrio", 1), ("Redigir a seção financeira", 2)],
 10: [("Montar a matriz SWOT", 1.5), ("Cruzar quadrantes (estratégias)", 1)],
}

def slots(livres, ate):
    """livres: {weekday: [(h_ini, h_fim)]}; gera (data, ini, fim) de HOJE até 'ate'."""
    d = HOJE
    while d <= ate:
        if d != FERIADO:
            for a, b in livres.get(d.weekday(), []):
                yield [d, a, b]
        d += dt.timedelta(days=1)

def semana(d):
    return 0 if d <= D(18,10) else 1 + (d - D(19,10)).days // 7

def planejar(p):
    perfil_p = p["tipo"][3] == "P"
    fator = 1.0 if perfil_p else 0.8           # J: 20% de folga
    folga_p = 2 if perfil_p else 0               # P: prazos internos 2 dias antes
    disp = [s for s in slots(p["livres"], D(11,11))]
    cap = {id(s): (s[2]-s[1]) * fator for s in disp}
    linhas, avisos, n = [], [], 0
    for b in sorted(p["blocos"], key=lambda x: MARCO[x]):
        limite = MARCO[b] - dt.timedelta(days=folga_p)
        pos = 0
        for nome, h in TAREFAS[b] + [("Revisar e marcar como concluído", 0.5)]:
            assert h <= 2, "tarefa > 2h viola a regra 1"
            s = next((s for s in disp[pos:] if s[0] <= limite and cap[id(s)] >= h), None)
            if s: pos = disp.index(s)
            if not s:
                avisos.append(f"Bloco {b}: sem horário para '{nome}' até {limite:%d/%m}")
                continue
            ini = round((s[1] + (s[2]-s[1])*fator - cap[id(s)]) * 2) / 2
            cap[id(s)] -= h
            n += 1
            prio = "Faça agora" if (limite - s[0]).days <= 2 else "Agende"
            linhas.append([n, semana(s[0]), s[0], DIAS[s[0].weekday()],
                           f"{int(ini)}h{int(ini%1*60):02d}–{int(ini+h)}h{int((ini+h)%1*60):02d}",
                           b, nome, p["nome"], h, prio, "No prazo", ""])
    # revisões de domingo e check-ins
    d = HOJE
    while d <= D(13,11):
        if d.weekday() == 6:
            linhas.append([None, semana(d), d, "dom", "20h00–20h15", "—", "Revisão semanal (15 min)", p["nome"], 0.25, "Agende", "No prazo", ""])
        if d.weekday() == p["checkin"][0]:
            linhas.append([None, semana(d), d, DIAS[d.weekday()], p["checkin"][1], "—", "Check-in do grupo (20 min)", "Grupo", 0.33, "Agende", "No prazo", ""])
        d += dt.timedelta(days=1)
    linhas.sort(key=lambda r: (r[2], r[4].zfill(12)))
    for i, r in enumerate(linhas, 1):
        r[0] = i
    # validações das regras
    seq = {}
    for r in linhas:
        if isinstance(r[5], int):
            assert (r[2], r[4].zfill(12)) >= seq.get(r[5], (HOJE, "")), "ordem do bloco violada"
            seq[r[5]] = (r[2], r[4].zfill(12))
        assert r[2] != FERIADO, "tarefa em feriado"
        if isinstance(r[5], int):
            assert r[2] <= MARCO[r[5]], "bloco termina depois do marco"
    return linhas, avisos

def excel(p, linhas, avisos, out):
    wb = Workbook()
    hdr = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="B65F3A")
    ws = wb.active; ws.title = "Cronograma"
    cols = ["ID","Semana","Data (dd/mm)","Dia","Horário","Bloco","Tarefa","Dona","Tempo (h)","Prioridade","Status","Observações"]
    ws.append(cols)
    for r in linhas:
        ws.append([r[0], r[1], r[2].strftime("%d/%m"), *r[3:]])
    for c in ws[1]: c.font, c.fill = hdr, fill
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"No prazo,Em risco,Atrasado,Concluído"'); ws.add_data_validation(dv)
    rng = f"K2:K{ws.max_row}"; dv.add(rng)
    for txt, cor in [("No prazo","C8E6C9"),("Em risco","FFE082"),("Atrasado","EF9A9A"),("Concluído","E0E0E0")]:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", fgColor=cor)))
    for col, w in zip("ABCDEFGHIJKL", [5,8,12,6,13,7,44,10,9,12,12,30]): ws.column_dimensions[col].width = w
    m = wb.create_sheet("Marcos")
    m.append(["Data","Marco","Blocos","% esperado","Blocos concluídos","% real","Situação"])
    for i,(d,mc,bl,pc) in enumerate([("18/10","Plano do grupo pronto","—",0),("25/10","Fundamentos","2, 3, 4",27),("01/11","Operação","5, 6, 7",55),("08/11","Números","8, 9, 10",82),("11/11","Versão final","11, 1",100),("13/11","Entrega","—",100)], start=2):
        m.append([d, mc, bl, pc/100, None, f"=IF(E{i}=\"\",\"\",E{i}/11)", f'=IF(E{i}="","",IF(F{i}>=D{i},"No ritmo","Atrás do marco"))'])
        m[f"D{i}"].number_format = m[f"F{i}"].number_format = "0%"
    for c in m[1]: c.font, c.fill = hdr, fill
    s = wb.create_sheet("Minha semana")
    s.append(["Horário"] + ["Segunda","Terça","Quarta","Quinta","Sexta","Sábado","Domingo"])
    for h in range(8, 23):
        row = [f"{h}h"]
        for wd in range(7):
            v = ""
            for a, b in p["livres"].get(wd, []):
                if a <= h < b: v = p["bloco_rotulo"]
            if wd == p["checkin"][0] and p["checkin"][1].startswith(f"{h}h"): v = "Check-in do grupo"
            if wd == 6 and h == 20: v = "Revisão semanal"
            row.append(v)
        s.append(row)
    for c in s[1]: c.font, c.fill = hdr, fill
    k = wb.create_sheet("Check-in")
    k.append(["Data","O que foi concluído","O que travou","Semáforo do marco","Tarefas e donas da próxima semana"])
    for c in k[1]: c.font, c.fill = hdr, fill
    wb.save(out)

ALUNAS = [
 dict(nome="Bia", tipo="ENFP", blocos=[3, 7, 8], checkin=(3, "18h00–18h20"),
      livres={0: [(19, 21)], 2: [(19, 21)], 5: [(9, 12)]}, bloco_rotulo="Projeto · tema do dia"),
 dict(nome="Carol", tipo="ISTJ", blocos=[2, 4, 10], checkin=(3, "18h00–18h20"),
      livres={1: [(19, 21)], 3: [(14, 17)], 5: [(14, 17)], 6: [(9, 11)]}, bloco_rotulo="Projeto · time blocking"),
]
out_dir = pathlib.Path(sys.argv[1]); out_dir.mkdir(parents=True, exist_ok=True)
for p in ALUNAS:
    linhas, avisos = planejar(p)
    f = out_dir / f"exemplo-cronograma-{p['nome'].lower()}-{p['tipo']}.xlsx"
    excel(p, linhas, avisos, f)
    tarefas = [r for r in linhas if isinstance(r[5], int)]
    print(f"\n== {p['nome']} ({p['tipo']}) · blocos {p['blocos']} · {len(tarefas)} tarefas agendadas · {f.name}")
    for r in tarefas[:4]: print("  ", r[2].strftime("%d/%m"), r[3], r[4], "| B", r[5], r[6])
    print("   ...")
    print("   AVISOS:" if avisos else "   Sem avisos: cabe nos horários.", *avisos, sep="\n    ")
